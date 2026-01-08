import csv
from openpyxl import load_workbook
import logging
import os

logging.basicConfig(level=logging.INFO)
LOGGER = logging.getLogger(__name__)

def excel_to_csv_all_sheets(excel_file_path, output_dir=".", prefix="output"):
    """Convert all sheets in an Excel workbook to separate CSV files.

    Args:
        excel_file_path (str | os.PathLike): Path to the .xlsx file.
        output_dir (str | os.PathLike): Directory where CSVs will be written.
        prefix (str): Prefix for generated CSV filenames.

    Returns:
        list[str]: List of generated CSV file paths.
    """
    os.makedirs(output_dir, exist_ok=True)
    wb = load_workbook(excel_file_path, read_only=True, data_only=True)
    generated = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))

        if not rows:
            LOGGER.warning("Sheet '%s' is empty, skipping...", sheet)
            continue

        headers = [str(h) if h is not None else "" for h in rows[0]]

        # Handle duplicate headers
        seen = {}
        for i, h in enumerate(headers):
            if h in seen:
                seen[h] += 1
                headers[i] = f"{h}_{seen[h]}"
            else:
                seen[h] = 0

        # Output CSV file per sheet
        safe_sheet_name = sheet.replace(" ", "_")
        csv_file_path = os.path.join(output_dir, f"{prefix}_{safe_sheet_name}.csv")

        with open(csv_file_path, "w", newline="", encoding="utf-8") as csvfile:
            writer = csv.writer(csvfile)
            writer.writerow(headers)
            for row in rows[1:]:
                writer.writerow(row)

        LOGGER.info("Sheet '%s' converted to '%s'", sheet, csv_file_path)
        generated.append(csv_file_path)
    return generated
