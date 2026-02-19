import unittest
import tempfile
import os
from openpyxl import Workbook
from openpyxl.comments import Comment
from datetime import datetime, date, time
from singer_encodings.excel_reader import get_excel_row_iterator
from singer_encodings.excel_helper import ExcelHelper


class TestBasicExcelReading(unittest.TestCase):
    """Test basic Excel file reading functionality."""
    def setUp(self):
        self.temp_file = tempfile.NamedTemporaryFile(mode='wb', suffix='.xlsx', delete=False)
        self.temp_file.close()
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append(["id", "name", "email", "age"])
        ws.append([1, "John Doe", "john@example.com", 30])
        ws.append([2, "Jane Smith", "jane@example.com", 25])
        ws.append([3, "Bob Johnson", "bob@example.com", 35])
        wb.save(self.temp_file.name)
        wb.close()

    def tearDown(self):
        if os.path.exists(self.temp_file.name):
            os.unlink(self.temp_file.name)

    def test_basic_reading(self):
        with open(self.temp_file.name, 'rb') as f:
            row_iterator = get_excel_row_iterator(f)
            self.assertIsNotNone(row_iterator)
            rows = list(row_iterator)
            self.assertEqual(len(rows), 3)
            sheet_name, first_row = rows[0]
            self.assertEqual(sheet_name, "Sheet1")
            self.assertEqual(first_row["id"], 1)
            self.assertEqual(first_row["name"], "John Doe")
            self.assertEqual(first_row["email"], "john@example.com")
            self.assertEqual(first_row["age"], 30)


class TestEmptyExcelFile(unittest.TestCase):
    """Test handling of empty Excel files."""
    def setUp(self):
        self.temp_file = tempfile.NamedTemporaryFile(mode='wb', suffix='.xlsx', delete=False)
        self.temp_file.close()
        wb = Workbook()
        ws = wb.active
        wb.save(self.temp_file.name)
        wb.close()

    def tearDown(self):
        if os.path.exists(self.temp_file.name):
            os.unlink(self.temp_file.name)

    def test_empty_file(self):
        with open(self.temp_file.name, 'rb') as f:
            row_iterator = get_excel_row_iterator(f)
            self.assertIsNone(row_iterator)


class TestMultipleSheets(unittest.TestCase):
    """Test reading Excel files with multiple sheets."""
    def setUp(self):
        self.temp_file = tempfile.NamedTemporaryFile(mode='wb', suffix='.xlsx', delete=False)
        self.temp_file.close()
        wb = Workbook()
        ws1 = wb.active
        ws1.title = "Employees"
        ws1.append(["id", "name"])
        ws1.append([1, "Alice"])
        ws1.append([2, "Bob"])
        ws2 = wb.create_sheet("Departments")
        ws2.append(["dept_id", "dept_name"])
        ws2.append([10, "Engineering"])
        ws2.append([20, "Sales"])
        wb.save(self.temp_file.name)
        wb.close()

    def tearDown(self):
        if os.path.exists(self.temp_file.name):
            os.unlink(self.temp_file.name)

    def test_all_sheets(self):
        with open(self.temp_file.name, 'rb') as f:
            row_iterator = get_excel_row_iterator(f)
            rows = list(row_iterator)
            self.assertEqual(len(rows), 4)
            self.assertEqual(rows[0][0], "Employees")
            self.assertEqual(rows[0][1]["id"], 1)
            self.assertEqual(rows[2][0], "Departments")
            self.assertEqual(rows[2][1]["dept_id"], 10)

    def test_specific_sheet(self):
        with open(self.temp_file.name, 'rb') as f:
            options = {"sheet_name": "Departments"}
            row_iterator = get_excel_row_iterator(f, options=options)
            rows = list(row_iterator)
            self.assertEqual(len(rows), 2)
            self.assertEqual(rows[0][0], "Departments")
            self.assertEqual(rows[0][1]["dept_id"], 10)


class TestDuplicateHeaders(unittest.TestCase):
    """Test handling of duplicate column headers."""
    def setUp(self):
        self.temp_file = tempfile.NamedTemporaryFile(mode='wb', suffix='.xlsx', delete=False)
        self.temp_file.close()
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append(["id", "name", "email", "name"])
        ws.append([1, "John", "john@example.com", "Johnny"])
        wb.save(self.temp_file.name)
        wb.close()

    def tearDown(self):
        if os.path.exists(self.temp_file.name):
            os.unlink(self.temp_file.name)

    def test_duplicate_headers(self):
        with open(self.temp_file.name, 'rb') as f:
            row_iterator = get_excel_row_iterator(f)
            rows = list(row_iterator)
            sheet_name, row = rows[0]
            self.assertEqual(row["name"], "John")
            self.assertIn("_sdc_extra", row)
            self.assertEqual(row["_sdc_extra"], [{"name": "Johnny"}])


class TestKeyPropertiesValidation(unittest.TestCase):
    """Test validation of required key properties."""
    def setUp(self):
        self.temp_file = tempfile.NamedTemporaryFile(mode='wb', suffix='.xlsx', delete=False)
        self.temp_file.close()
        wb = Workbook()
        ws = wb.active
        ws.append(["id", "name", "email"])
        ws.append([1, "John", "john@example.com"])
        wb.save(self.temp_file.name)
        wb.close()

    def tearDown(self):
        if os.path.exists(self.temp_file.name):
            os.unlink(self.temp_file.name)

    def test_valid_key_properties(self):
        with open(self.temp_file.name, 'rb') as f:
            options = {"key_properties": ["id", "email"]}
            row_iterator = get_excel_row_iterator(f, options=options)
            rows = list(row_iterator)
            self.assertEqual(len(rows), 1)

    def test_missing_key_properties(self):
        with open(self.temp_file.name, 'rb') as f:
            options = {"key_properties": ["id", "missing_column"]}
            with self.assertRaises(ValueError) as context:
                row_iterator = get_excel_row_iterator(f, options=options)
                list(row_iterator)
            self.assertIn("missing required headers", str(context.exception))


class TestDateOverridesValidation(unittest.TestCase):
    """Test validation of date override columns."""
    def setUp(self):
        self.temp_file = tempfile.NamedTemporaryFile(mode='wb', suffix='.xlsx', delete=False)
        self.temp_file.close()
        wb = Workbook()
        ws = wb.active
        ws.append(["id", "created_at", "updated_at"])
        ws.append([1, datetime(2024, 1, 15, 10, 30), datetime(2024, 2, 20, 15, 45)])
        wb.save(self.temp_file.name)
        wb.close()

    def tearDown(self):
        if os.path.exists(self.temp_file.name):
            os.unlink(self.temp_file.name)

    def test_valid_date_overrides(self):
        with open(self.temp_file.name, 'rb') as f:
            options = {"date_overrides": ["created_at", "updated_at"]}
            row_iterator = get_excel_row_iterator(f, options=options)
            rows = list(row_iterator)
            self.assertEqual(len(rows), 1)

    def test_missing_date_overrides(self):
        with open(self.temp_file.name, 'rb') as f:
            options = {"date_overrides": ["created_at", "deleted_at"]}
            with self.assertRaises(ValueError) as context:
                row_iterator = get_excel_row_iterator(f, options=options)
                list(row_iterator)
            self.assertIn("missing date_overrides headers", str(context.exception))


class TestDateTimeNormalization(unittest.TestCase):
    """Test automatic datetime normalization to ISO-8601."""
    def setUp(self):
        self.temp_file = tempfile.NamedTemporaryFile(mode='wb', suffix='.xlsx', delete=False)
        self.temp_file.close()
        wb = Workbook()
        ws = wb.active
        ws.append(["id", "timestamp", "birth_date", "start_time"])
        ws.append([1, datetime(2024, 1, 15, 14, 30, 0), date(1990, 5, 20), time(9, 30, 0)])
        wb.save(self.temp_file.name)
        wb.close()

    def tearDown(self):
        if os.path.exists(self.temp_file.name):
            os.unlink(self.temp_file.name)

    def test_datetime_normalization(self):
        with open(self.temp_file.name, 'rb') as f:
            row_iterator = get_excel_row_iterator(f)
            rows = list(row_iterator)
            sheet_name, row = rows[0]
            self.assertEqual(row["timestamp"], "2024-01-15T14:30:00Z")
            self.assertEqual(row["birth_date"], "1990-05-20T00:00:00Z")
            self.assertEqual(row["start_time"], "09:30:00")


class TestHeadersInCatalog(unittest.TestCase):
    """Test filtering columns based on catalog headers."""
    def setUp(self):
        self.temp_file = tempfile.NamedTemporaryFile(mode='wb', suffix='.xlsx', delete=False)
        self.temp_file.close()
        wb = Workbook()
        ws = wb.active
        ws.append(["id", "name", "extra_field", "another_extra"])
        ws.append([1, "John", "extra_value", "another_value"])
        wb.save(self.temp_file.name)
        wb.close()

    def tearDown(self):
        if os.path.exists(self.temp_file.name):
            os.unlink(self.temp_file.name)

    def test_headers_in_catalog(self):
        with open(self.temp_file.name, 'rb') as f:
            headers_in_catalog = ["id", "name"]
            row_iterator = get_excel_row_iterator(f, headers_in_catalog=headers_in_catalog)
            rows = list(row_iterator)
            sheet_name, row = rows[0]
            self.assertEqual(row["id"], 1)
            self.assertEqual(row["name"], "John")
            self.assertIn("_sdc_extra", row)
            extra = row["_sdc_extra"]
            self.assertEqual(len(extra), 2)


class TestHyperlinkPreservation(unittest.TestCase):
    """Test preservation of cell hyperlinks."""
    def setUp(self):
        self.temp_file = tempfile.NamedTemporaryFile(mode='wb', suffix='.xlsx', delete=False)
        self.temp_file.close()
        wb = Workbook()
        ws = wb.active
        ws.append(["id", "website"])
        ws.append([1, "Visit Site"])
        ws['B2'].hyperlink = "https://example.com"
        wb.save(self.temp_file.name)
        wb.close()

    def tearDown(self):
        if os.path.exists(self.temp_file.name):
            os.unlink(self.temp_file.name)

    def test_hyperlink_preservation(self):
        with open(self.temp_file.name, 'rb') as f:
            row_iterator = get_excel_row_iterator(f)
            rows = list(row_iterator)
            sheet_name, row = rows[0]
            self.assertIsInstance(row["website"], list)
            self.assertEqual(len(row["website"]), 1)
            self.assertEqual(row["website"][0]["text"], "Visit Site")
            self.assertEqual(row["website"][0]["url"], "https://example.com")


class TestCommentExtraction(unittest.TestCase):
    """Test extraction of cell comments."""
    def setUp(self):
        self.temp_file = tempfile.NamedTemporaryFile(mode='wb', suffix='.xlsx', delete=False)
        self.temp_file.close()
        wb = Workbook()
        ws = wb.active
        ws.append(["id", "notes"])
        ws.append([1, "Important data"])
        comment = Comment("Review this value", "John Doe")
        ws['B2'].comment = comment
        wb.save(self.temp_file.name)
        wb.close()

    def tearDown(self):
        if os.path.exists(self.temp_file.name):
            os.unlink(self.temp_file.name)

    def test_comment_extraction(self):
        with open(self.temp_file.name, 'rb') as f:
            row_iterator = get_excel_row_iterator(f)
            rows = list(row_iterator)
            sheet_name, row = rows[0]
            self.assertIsInstance(row["notes"], list)
            self.assertEqual(len(row["notes"]), 1)
            self.assertEqual(row["notes"][0]["text"], "Important data")
            self.assertIn("comment", row["notes"][0])


class TestExcelHelper(unittest.TestCase):
    """Test ExcelHelper utility methods."""
    def test_convert_datetime_to_iso(self):
        helper = ExcelHelper()
        dt = datetime(2024, 1, 15, 14, 30, 0)
        result = helper.convert_for_json(dt)
        self.assertEqual(result, "2024-01-15T14:30:00Z")

    def test_convert_date_to_iso(self):
        helper = ExcelHelper()
        d = date(2024, 1, 15)
        result = helper.convert_for_json(d)
        self.assertEqual(result, "2024-01-15T00:00:00Z")

    def test_convert_time_to_string(self):
        helper = ExcelHelper()
        t = time(14, 30, 0)
        result = helper.convert_for_json(t)
        self.assertEqual(result, "14:30:00")

    def test_convert_date_string_formats(self):
        helper = ExcelHelper()
        test_cases = [
            ("2024-01-15", "2024-01-15T00:00:00Z"),
            ("15-Jan-24", "2024-01-15T00:00:00Z"),
            ("01/15/2024", "2024-01-15T00:00:00Z"),
        ]
        for input_str, expected in test_cases:
            result = helper.convert_for_json(input_str)
            self.assertEqual(result, expected, f"Failed for input: {input_str}")


class TestSkipEmptyRows(unittest.TestCase):
    """Test skipping completely empty rows."""
    def setUp(self):
        self.temp_file = tempfile.NamedTemporaryFile(mode='wb', suffix='.xlsx', delete=False)
        self.temp_file.close()
        wb = Workbook()
        ws = wb.active
        ws.append(["id", "name"])
        ws.append([1, "John"])
        ws.append([None, None])
        ws.append([2, "Jane"])
        ws.append([None, None])
        wb.save(self.temp_file.name)
        wb.close()

    def tearDown(self):
        if os.path.exists(self.temp_file.name):
            os.unlink(self.temp_file.name)

    def test_skip_empty_rows(self):
        with open(self.temp_file.name, 'rb') as f:
            row_iterator = get_excel_row_iterator(f)
            rows = list(row_iterator)
            self.assertEqual(len(rows), 2)
            self.assertEqual(rows[0][1]["id"], 1)
            self.assertEqual(rows[1][1]["id"], 2)


class TestInvalidSheetName(unittest.TestCase):
    """Test handling of invalid sheet names."""
    def setUp(self):
        self.temp_file = tempfile.NamedTemporaryFile(mode='wb', suffix='.xlsx', delete=False)
        self.temp_file.close()
        wb = Workbook()
        ws = wb.active
        ws.title = "ValidSheet"
        ws.append(["id", "name"])
        ws.append([1, "John"])
        wb.save(self.temp_file.name)
        wb.close()

    def tearDown(self):
        if os.path.exists(self.temp_file.name):
            os.unlink(self.temp_file.name)

    def test_invalid_sheet_name(self):
        with open(self.temp_file.name, 'rb') as f:
            options = {"sheet_name": "NonExistentSheet"}
            with self.assertRaises(ValueError) as context:
                row_iterator = get_excel_row_iterator(f, options=options)
                if row_iterator:
                    list(row_iterator)
            self.assertIn("not found", str(context.exception))
            self.assertIn("Available sheets", str(context.exception))


class TestWorkbookCleanupOnException(unittest.TestCase):
    """Test that workbook is properly closed even when exceptions occur."""
    def setUp(self):
        self.temp_file = tempfile.NamedTemporaryFile(mode='wb', suffix='.xlsx', delete=False)
        self.temp_file.close()
        wb = Workbook()
        ws = wb.active
        ws.append(["id", "name"])
        ws.append([1, "Test"])
        wb.save(self.temp_file.name)
        wb.close()

    def tearDown(self):
        if os.path.exists(self.temp_file.name):
            os.unlink(self.temp_file.name)

    def test_workbook_cleanup_on_invalid_sheet(self):
        """Verify workbook is closed when invalid sheet name is requested."""
        with open(self.temp_file.name, 'rb') as f:
            options = {'sheet_name': 'NonExistent'}
            with self.assertRaises(ValueError):
                row_iterator = get_excel_row_iterator(f, options=options)
                if row_iterator:
                    list(row_iterator)
        with open(self.temp_file.name, 'rb') as f:
            row_iterator = get_excel_row_iterator(f)
            rows = list(row_iterator)
            self.assertEqual(len(rows), 1)

    def test_workbook_cleanup_on_validation_error(self):
        """Verify workbook is closed when validation fails."""
        with open(self.temp_file.name, 'rb') as f:
            options = {'key_properties': ['nonexistent_column']}
            with self.assertRaises(ValueError):
                row_iterator = get_excel_row_iterator(f, options=options)
                if row_iterator:
                    list(row_iterator)
        with open(self.temp_file.name, 'rb') as f:
            row_iterator = get_excel_row_iterator(f)
            rows = list(row_iterator)
            self.assertEqual(len(rows), 1)


class TestExcessValues(unittest.TestCase):
    """Test handling of rows with more values than headers."""
    def setUp(self):
        self.temp_file = tempfile.NamedTemporaryFile(mode='wb', suffix='.xlsx', delete=False)
        self.temp_file.close()
        wb = Workbook()
        ws = wb.active
        ws.append(["id", "name", "email"])
        ws.append([1, "John Doe", "john@example.com", "Extra1", "Extra2"])
        ws.append([2, "Jane Smith", "jane@example.com", "Extra3"])
        ws.append([3, "Bob Johnson", "bob@example.com"])
        wb.save(self.temp_file.name)
        wb.close()

    def tearDown(self):
        if os.path.exists(self.temp_file.name):
            os.unlink(self.temp_file.name)

    def test_excess_values_handling(self):
        """Test that excess values beyond header count are stored in _sdc_extra.

        When a row has more values than the original headers, Excel reads those
        additional columns as having empty header names. The first empty header
        is treated as a unique field, while subsequent empty headers trigger
        duplicate detection and are stored in _sdc_extra.
        """
        with open(self.temp_file.name, 'rb') as f:
            row_iterator = get_excel_row_iterator(f)
            rows = list(row_iterator)
            self.assertEqual(len(rows), 3)
            sheet_name, row1 = rows[0]
            self.assertEqual(row1["id"], 1)
            self.assertEqual(row1["name"], "John Doe")
            self.assertEqual(row1["email"], "john@example.com")
            self.assertEqual(row1[''], 'Extra1')
            self.assertIn("_sdc_extra", row1)
            self.assertEqual(len(row1["_sdc_extra"]), 1)
            self.assertEqual(row1["_sdc_extra"][0][''], 'Extra2')
            sheet_name, row2 = rows[1]
            self.assertEqual(row2["id"], 2)
            self.assertEqual(row2[''], 'Extra3')
            self.assertIn("_sdc_extra", row2)
            self.assertEqual(row2["_sdc_extra"][0][''], None)
            sheet_name, row3 = rows[2]
            self.assertEqual(row3["id"], 3)
            self.assertEqual(row3[''], None)
            self.assertIn("_sdc_extra", row3)
            self.assertEqual(row3["_sdc_extra"][0][''], None)


class TestHyperlinkAndComment(unittest.TestCase):
    """Test handling of cells with both hyperlinks and comments."""
    def setUp(self):
        self.temp_file = tempfile.NamedTemporaryFile(mode='wb', suffix='.xlsx', delete=False)
        self.temp_file.close()
        wb = Workbook()
        ws = wb.active
        ws.append(["id", "website"])
        ws.append([1, "Visit Site"])
        ws['B2'].hyperlink = "https://example.com"
        comment = Comment("Check this link", "Reviewer")
        ws['B2'].comment = comment
        wb.save(self.temp_file.name)
        wb.close()

    def tearDown(self):
        if os.path.exists(self.temp_file.name):
            os.unlink(self.temp_file.name)

    def test_hyperlink_and_comment_combination(self):
        """Test that cells with both hyperlinks and comments preserve both.

        When a cell has both a hyperlink and a comment, they are combined
        into a single structured dictionary containing all the data.
        """
        with open(self.temp_file.name, 'rb') as f:
            row_iterator = get_excel_row_iterator(f)
            rows = list(row_iterator)
            sheet_name, row = rows[0]
            self.assertIsInstance(row["website"], list)
            self.assertEqual(len(row["website"]), 1)
            cell_data = row["website"][0]
            self.assertEqual(cell_data["text"], "Visit Site")
            self.assertEqual(cell_data["url"], "https://example.com")
            self.assertIn("comment", cell_data)
            self.assertIsInstance(cell_data["comment"], dict)
            self.assertEqual(cell_data["comment"]["text"], "Check this link")
            self.assertEqual(cell_data["comment"]["excel_author"], "Reviewer")


if __name__ == '__main__':
    unittest.main()
